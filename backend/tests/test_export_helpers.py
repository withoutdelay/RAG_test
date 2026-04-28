from __future__ import annotations

import asyncio
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

from docx import Document as DocxDocument

from app.services.export.service import (
    ExportService,
    build_export_snapshot,
    render_export_markdown,
    render_export_sections_markdown,
    summarize_holistic_finalization_trace,
    write_export_file,
)


class _FakeHolistic:
    def __init__(
        self,
        *,
        content: str | None = None,
        contents: list[str | Exception] | None = None,
        error: Exception | None = None,
    ) -> None:
        self.content = content
        self.contents = contents or []
        self.error = error
        self.calls: list[dict] = []

    async def finalize(self, **kwargs):
        self.calls.append(kwargs)
        if self.error:
            raise self.error
        if self.contents:
            result = self.contents[min(len(self.calls) - 1, len(self.contents) - 1)]
            if isinstance(result, Exception):
                raise result
            return SimpleNamespace(content=result)
        return SimpleNamespace(content=self.content or "")


class ExportHelperTests(unittest.TestCase):
    def test_build_export_snapshot_contains_bound_versions(self) -> None:
        snapshot = build_export_snapshot(
            project=SimpleNamespace(id=uuid4(), status="EXPORTABLE", current_draft_version=2),
            outline=SimpleNamespace(id=uuid4(), version=3),
            requirement_card=SimpleNamespace(id=uuid4(), version=4),
            evidence_bundle=SimpleNamespace(id=uuid4(), retrieval_version=5),
            validation_report=SimpleNamespace(id=uuid4(), status="passed"),
            review_tasks=[
                SimpleNamespace(id=uuid4(), task_type="final_review", status="resolved", blocking_level="P0")
            ],
        )

        self.assertEqual(snapshot["draft_version"], 2)
        self.assertEqual(snapshot["outline_version"], 3)
        self.assertEqual(snapshot["requirement_card_version"], 4)
        self.assertEqual(snapshot["evidence_bundle_version"], 5)
        self.assertEqual(snapshot["validation_status"], "passed")
        self.assertEqual(snapshot["review_tasks"][0]["task_type"], "final_review")

    def test_render_export_markdown_keeps_sections_and_citations(self) -> None:
        markdown = render_export_markdown(
            project=SimpleNamespace(name="测试项目"),
            outline=SimpleNamespace(
                outline_json={
                    "title": "测试项目技术方案",
                    "sections": [
                        {"section_id": "1", "title": "项目概述", "children": []},
                        {"section_id": "2", "title": "技术架构", "children": []},
                    ],
                }
            ),
            section_drafts=[
                SimpleNamespace(
                    section_id="1",
                    title="项目概述",
                    content_md="## 项目概述\n\n说明项目背景。",
                    citation_refs=[],
                ),
                SimpleNamespace(
                    section_id="2",
                    title="技术架构",
                    content_md="## 技术架构\n\n说明系统架构。\n\n[[ASSET:FIGURE:asset-001]]",
                    citation_refs=[
                        {
                            "evidence_id": "ev_001",
                            "source_title": "历史方案A",
                            "heading_path": ["第2章", "技术架构"],
                        }
                    ],
                    validator_result={
                        "recommended_assets": [
                            {
                                "asset_id": "asset-001",
                                "asset_type": "figure",
                                "title": "系统拓扑图",
                                "document_name": "历史方案A",
                                "page_no": 12,
                                "heading_path": "第2章 > 技术架构",
                                "preview_text": "站控层、间隔层和网络层拓扑示意。",
                                "reason": "与章节《技术架构》高度相关，可作为当前方案的参考插图。",
                                "review_required": True,
                            }
                        ]
                    },
                ),
            ],
            snapshot={
                "exported_at": "2026-03-23T12:00:00+08:00",
                "draft_version": 2,
                "validation_report_id": str(uuid4()),
            },
        )

        self.assertIn("# 测试项目技术方案", markdown)
        self.assertIn("## 项目概述", markdown)
        self.assertIn("## 技术架构", markdown)
        self.assertIn("## 引用清单", markdown)
        self.assertIn("ev_001 / 历史方案A / 第2章 > 技术架构", markdown)
        self.assertIn("[建议插入图片] 系统拓扑图", markdown)
        self.assertIn("资产占位符：[[ASSET:FIGURE:asset-001]]", markdown)
        self.assertIn("来源：历史方案A / 第 12 页 / 第2章 > 技术架构", markdown)

    def test_render_export_sections_markdown_can_keep_raw_asset_placeholders(self) -> None:
        markdown = render_export_sections_markdown(
            outline=SimpleNamespace(
                outline_json={
                    "sections": [
                        {"section_id": "1", "title": "技术架构", "children": []},
                    ]
                }
            ),
            section_drafts=[
                SimpleNamespace(
                    section_id="1",
                    title="技术架构",
                    content_md="## 技术架构\n\n说明系统架构。\n\n[[ASSET:FIGURE:asset-001]]",
                    citation_refs=[],
                    validator_result={
                        "recommended_assets": [
                            {"asset_id": "asset-001", "asset_type": "figure", "title": "系统拓扑图"}
                        ]
                    },
                )
            ],
            render_assets=False,
        )

        self.assertIn("[[ASSET:FIGURE:asset-001]]", markdown)
        self.assertNotIn("[建议插入图片]", markdown)

    def test_write_export_file_creates_word_docx(self) -> None:
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as handle:
            path = Path(handle.name)
        try:
            write_export_file(
                markdown="# 测试方案\n\n## 技术参数\n\n| 参数 | 值 |\n| --- | --- |\n| 电压 | 10kV |\n\n- 支持联锁控制",
                file_format="docx",
                destination=path,
            )
            document = DocxDocument(str(path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("测试方案", text)
            self.assertIn("技术参数", text)
            self.assertEqual(document.tables[0].cell(1, 1).text, "10kV")
        finally:
            path.unlink(missing_ok=True)

    def test_export_holistic_finalization_is_optional_and_renders_assets(self) -> None:
        finalized = (
            "# 测试项目技术方案\n\n"
            "## 技术架构\n\n"
            + ("融合后正文保留系统架构、接口、联锁、供货边界和技术参数。" * 8)
            + "\n\n[[ASSET:FIGURE:asset-001]]\n"
        )
        holistic = _FakeHolistic(content=finalized)
        service = ExportService(
            settings=SimpleNamespace(export_holistic_finalization_enabled=True),
            holistic=holistic,
        )
        current_markdown = "# 测试项目技术方案\n\n## 技术架构\n\n原始正文。\n"
        markdown, trace = asyncio.run(
            service._maybe_finalize_markdown(
                task_id="export-test",
                project=SimpleNamespace(id=uuid4(), name="测试项目", current_draft_version=2),
                outline=SimpleNamespace(
                    outline_json={
                        "title": "测试项目技术方案",
                        "sections": [{"section_id": "1", "title": "技术架构", "children": []}],
                    }
                ),
                requirement_card=SimpleNamespace(content={"global_params": {"voltage_level": "10kV"}}),
                section_drafts=[
                    SimpleNamespace(
                        section_id="1",
                        title="技术架构",
                        content_md="## 技术架构\n\n" + ("原始正文。" * 30) + "\n\n[[ASSET:FIGURE:asset-001]]",
                        citation_refs=[
                            {
                                "evidence_id": "ev_001",
                                "source_title": "历史方案A",
                                "heading_path": ["第2章", "技术架构"],
                            }
                        ],
                        validator_result={
                            "recommended_assets": [
                                {
                                    "asset_id": "asset-001",
                                    "asset_type": "figure",
                                    "title": "系统拓扑图",
                                    "document_name": "历史方案A",
                                }
                            ]
                        },
                    )
                ],
                current_markdown=current_markdown,
            )
        )

        self.assertEqual(trace["status"], "succeeded")
        self.assertEqual(holistic.calls[0]["global_params"]["voltage_level"], "10kV")
        self.assertIn("融合后正文", markdown)
        self.assertIn("[建议插入图片] 系统拓扑图", markdown)
        self.assertIn("## 引用清单", markdown)

    def test_export_holistic_section_mode_falls_back_per_section(self) -> None:
        holistic = _FakeHolistic(
            contents=[
                "## 第一章\n\n## 第一章\n\n" + ("第一章融合后正文保留关键边界。" * 8),
                "太短",
            ]
        )
        service = ExportService(
            settings=SimpleNamespace(export_holistic_finalization_enabled=True, export_holistic_finalization_mode="section"),
            holistic=holistic,
        )
        current_markdown = "# 测试项目技术方案\n\n## 第一章\n\n原始一。\n\n## 第二章\n\n原始二。\n"
        markdown, trace = asyncio.run(
            service._maybe_finalize_markdown(
                task_id="export-test",
                project=SimpleNamespace(id=uuid4(), name="测试项目", current_draft_version=2),
                outline=SimpleNamespace(
                    outline_json={
                        "title": "测试项目技术方案",
                        "sections": [
                            {"section_id": "1", "title": "第一章", "children": []},
                            {"section_id": "2", "title": "第二章", "children": []},
                        ],
                    }
                ),
                requirement_card=SimpleNamespace(content={}),
                section_drafts=[
                    SimpleNamespace(
                        section_id="1",
                        title="第一章",
                        content_md="## 第一章\n\n" + ("原始一。" * 30),
                        citation_refs=[],
                        validator_result={},
                    ),
                    SimpleNamespace(
                        section_id="2",
                        title="第二章",
                        content_md="## 第二章\n\n" + ("原始二。" * 30),
                        citation_refs=[],
                        validator_result={},
                    ),
                ],
                current_markdown=current_markdown,
            )
        )

        self.assertEqual(trace["mode"], "section")
        self.assertEqual(trace["status"], "partial_fallback")
        self.assertEqual([item["status"] for item in trace["sections"]], ["succeeded", "fallback"])
        self.assertIn("第一章融合后正文", markdown)
        self.assertNotIn("## 第一章\n\n## 第一章", markdown)
        self.assertIn("## 第二章", markdown)
        self.assertEqual(len(holistic.calls), 2)

    def test_export_holistic_finalization_falls_back_when_output_is_unsafe(self) -> None:
        service = ExportService(
            settings=SimpleNamespace(export_holistic_finalization_enabled=True),
            holistic=_FakeHolistic(content="太短"),
        )
        current_markdown = "# 测试项目技术方案\n\n## 技术架构\n\n原始正文。\n"
        markdown, trace = asyncio.run(
            service._maybe_finalize_markdown(
                task_id="export-test",
                project=SimpleNamespace(id=uuid4(), name="测试项目", current_draft_version=2),
                outline=SimpleNamespace(
                    outline_json={
                        "title": "测试项目技术方案",
                        "sections": [{"section_id": "1", "title": "技术架构", "children": []}],
                    }
                ),
                requirement_card=SimpleNamespace(content={}),
                section_drafts=[
                    SimpleNamespace(
                        section_id="1",
                        title="技术架构",
                        content_md="## 技术架构\n\n" + ("原始正文。" * 30),
                        citation_refs=[],
                        validator_result={},
                    )
                ],
                current_markdown=current_markdown,
            )
        )

        self.assertEqual(markdown, current_markdown)
        self.assertEqual(trace["status"], "fallback")

    def test_summarize_holistic_finalization_trace_counts_section_fallbacks(self) -> None:
        summary = summarize_holistic_finalization_trace(
            {
                "enabled": True,
                "mode": "section",
                "status": "partial_fallback",
                "input_chars": 1200,
                "output_chars": 1500,
                "succeeded_sections": 2,
                "sections": [
                    {"section_id": "1", "status": "succeeded"},
                    {"section_id": "2", "status": "fallback"},
                    {"section_id": "3", "status": "succeeded"},
                ],
            }
        )

        self.assertEqual(
            summary,
            {
                "enabled": True,
                "mode": "section",
                "status": "partial_fallback",
                "input_chars": 1200,
                "output_chars": 1500,
                "section_count": 3,
                "succeeded_sections": 2,
                "fallback_sections": 1,
            },
        )


if __name__ == "__main__":
    unittest.main()
